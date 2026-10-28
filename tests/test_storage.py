"""Saving recordings, privacy, encryption, export and merging."""

import json
import subprocess
import sys
import zipfile
from pathlib import Path

import soundfile as sf
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric import rsa
from openpyxl import load_workbook

from recorder import crypto, ids
from recorder.audio import AudioSettings, save_asr_copy, save_recording
from recorder.export import export_speaker
from recorder.storage import SpeakerStore, list_speakers
from tests.conftest import speech_like

ROOT = Path(__file__).resolve().parent.parent
PRIVATE = {"name": "Ravi Kumar", "mobile": "9876543210", "email": "ravi@example.com",
           "pincode": "751001"}


def make_keys(folder: Path) -> tuple[Path, Path]:
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    pub, priv = folder / "public_key.pem", folder / "private_key.pem"
    pub.write_bytes(key.public_key().public_bytes(
        serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo))
    priv.write_bytes(key.private_bytes(serialization.Encoding.PEM,
                                       serialization.PrivateFormat.PKCS8,
                                       serialization.NoEncryption()))
    return pub, priv


def record_one(store: SpeakerStore, number: int, text: str) -> str:
    settings = AudioSettings(sample_rate=48000, bit_depth="24-bit PCM", channels=1,
                             file_format="WAV", asr_copy=True)
    take = speech_like()
    sid = ids.script_sentence_id(store.speaker_id, number)
    path, asr = store.audio_path(sid, "wav"), store.asr_path(sid)
    save_recording(take, settings, path)
    save_asr_copy(take, 48000, asr)
    store.add_record({"sentence_id": sid, "speaker_id": store.speaker_id, "text": text,
                      "audio_filepath": path.relative_to(store.dir).as_posix(),
                      "audio_16k_filepath": asr.relative_to(store.dir).as_posix(),
                      "duration_sec": 3.0})
    return sid


def new_store(root: Path) -> SpeakerStore:
    sid = "eng_OD_KHO_M_1999_" + ids.speaker_fingerprint("Ravi Kumar", 1999, "M", "751001",
                                                         "9876543210")
    store = SpeakerStore(root, sid)
    store.save_profile({"speaker_id": sid, "gender": "Male"}, PRIVATE)
    store.add_consent({"agreed": True, "agreed_at": "2026-10-28T10:00:00+05:30"})
    return store


def test_audio_files_and_transcripts(tmp_path):
    store = new_store(tmp_path / "rec")
    sid = record_one(store, 1, "Hello world.")
    master, asr = sf.info(str(store.audio_path(sid, "wav"))), sf.info(str(store.asr_path(sid)))
    assert (master.samplerate, master.subtype) == (48000, "PCM_24")
    assert (asr.samplerate, asr.subtype, asr.channels) == (16000, "PCM_16", 1)
    rows = list(load_workbook(store.dir / "transcripts.xlsx").active.values)
    assert rows == [("sentence_id", "sentence"), (sid, "Hello world.")]
    assert SpeakerStore(tmp_path / "rec", store.speaker_id).has(sid)   # reloads from disk


def test_rerecording_replaces_entry(tmp_path):
    store = new_store(tmp_path / "rec")
    record_one(store, 1, "First take.")
    record_one(store, 1, "Second take.")
    assert len(store.records) == 1
    assert list(store.records.values())[0]["text"] == "Second take."


def test_personal_details_stay_private(tmp_path):
    store = new_store(tmp_path / "rec")
    record_one(store, 1, "Hello world.")
    public = "".join((store.dir / f).read_text(encoding="utf-8-sig")
                     for f in ("speaker.json", "manifest.jsonl", "manifest.csv", "consent.json"))
    for value in PRIVATE.values():
        assert value.lower() not in public.lower()
    assert (store.dir / "private" / "contact.json").exists()


def test_contact_details_encrypted_with_public_key(tmp_path, monkeypatch):
    pub, priv = make_keys(tmp_path)
    monkeypatch.setattr(crypto, "PUBLIC_KEY_FILE", pub)
    store = new_store(tmp_path / "rec")
    enc = store.dir / "private" / "contact.enc"
    assert enc.exists() and not (store.dir / "private" / "contact.json").exists()
    assert "Ravi" not in enc.read_text()
    assert crypto.decrypt_json(json.loads(enc.read_text()), priv)["name"] == "Ravi Kumar"


def test_remove_if_empty(tmp_path):
    store = new_store(tmp_path / "rec")
    assert store.remove_if_empty() and not store.dir.exists()
    store = new_store(tmp_path / "rec")
    record_one(store, 1, "Keep me.")
    assert not store.remove_if_empty() and store.dir.exists()


def test_export_and_merge(tmp_path):
    pub, priv = make_keys(tmp_path)
    store = new_store(tmp_path / "rec")
    record_one(store, 1, "Hello world.")
    record_one(store, 2, "Second sentence.")
    assert list_speakers(tmp_path / "rec")[0]["clips"] == 2

    zip_path = export_speaker(store.dir, tmp_path / "subs")
    names = zipfile.ZipFile(zip_path).namelist()
    assert f"{store.speaker_id}/checksums.sha256" in names
    assert f"{store.speaker_id}/transcripts.xlsx" in names

    result = subprocess.run([sys.executable, str(ROOT / "tools" / "merge_submissions.py"),
                             str(tmp_path / "subs"), str(tmp_path / "merged")],
                            capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    merged = tmp_path / "merged"
    report = (merged / "merge_report.txt").read_text(encoding="utf-8")
    assert "Clips: 2" in report and "CHECKSUM FAIL" not in report
    rows = list(load_workbook(merged / "transcripts.xlsx").active.values)
    assert len(rows) == 3
    assert (merged / "PRIVATE_contacts.xlsx").exists()
    assert not (merged / "speakers" / store.speaker_id / "private").exists()
