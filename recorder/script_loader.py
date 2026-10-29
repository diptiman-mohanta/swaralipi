"""Loads a reading script (xlsx or csv) of sentence numbers and sentences."""

import csv
from pathlib import Path

from openpyxl import load_workbook

from .ids import format_script_number, normalize_text

NUMBER_HEADERS = {"sentence_no", "sentence_number", "sentence_num", "sentence_id", "id", "no",
                  "number", "sl_no", "slno", "s_no", "sno", "sr_no", "srno", "serial", "index"}
TEXT_HEADERS = {"sentence", "sentences", "text", "transcript", "transcription", "prompt",
                "utterance"}


class ScriptError(Exception):
    pass


def _key(h) -> str:
    return str(h or "").strip().lower().replace(" ", "_").replace(".", "").replace("-", "_")


def _read_rows(path: Path) -> list[list]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
        return rows
    if path.suffix.lower() in (".csv", ".tsv", ".txt"):
        raw = path.read_bytes()
        # Excel's "Unicode Text" export is UTF-16; Indian-language scripts need Unicode.
        encodings = ("utf-16",) if raw[:2] in (b"\xff\xfe", b"\xfe\xff") \
            else ("utf-8-sig", "cp1252")
        for enc in encodings:
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ScriptError("Could not read the file - save it as UTF-8 CSV.")
        # The delimiter is whichever of , ; or tab the first line uses most (csv.Sniffer
        # guesses wrong on some files, e.g. Hindi text with Windows line endings).
        lines = text.splitlines()
        first = next((line for line in lines if line.strip()), "")
        delimiter = max(",;\t", key=first.count) if any(d in first for d in ",;\t") else ","
        return [r for r in csv.reader(lines, delimiter=delimiter)]
    raise ScriptError("Unsupported file type - use .xlsx or .csv")


def load_script(path) -> list[tuple[str, str]]:
    """Returns [(sentence_number, sentence_text)] in file order."""
    path = Path(path)
    rows = [r for r in _read_rows(path) if any(str(c or "").strip() for c in r)]
    if not rows:
        raise ScriptError("The file is empty.")

    header = [_key(c) for c in rows[0]]
    num_col = next((i for i, h in enumerate(header) if h in NUMBER_HEADERS), None)
    text_col = next((i for i, h in enumerate(header) if h in TEXT_HEADERS), None)
    if text_col is not None:
        rows = rows[1:]
    else:
        # No recognised header: column 1 = number, column 2 = sentence (or text only).
        width = max(len(r) for r in rows)
        num_col, text_col = (0, 1) if width >= 2 else (None, 0)
        first = str(rows[0][num_col] if num_col is not None else "").strip()
        if num_col is not None and not format_script_number(first).isdigit():
            rows = rows[1:]   # unrecognised header row

    out, seen = [], set()
    for i, r in enumerate(rows, start=1):
        text = normalize_text(str(r[text_col] or "")) if text_col < len(r) else ""
        if not text:
            continue
        raw_num = r[num_col] if num_col is not None and num_col < len(r) else None
        number = str(raw_num).strip() if raw_num not in (None, "") else str(i)
        if not format_script_number(number):
            raise ScriptError(f"Row {i}: sentence number '{number}' is not valid.")
        key = format_script_number(number)
        if key in seen:
            raise ScriptError(f"Sentence number '{number}' appears more than once.")
        seen.add(key)
        out.append((number, text))
    if not out:
        raise ScriptError("No sentences found.")
    return out
